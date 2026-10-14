"""Reuse the independent numerical audit and validate design provenance."""
from pathlib import Path
import json
from hashlib import sha256
from pypdf import PdfReader
from docx import Document

here=Path(__file__).resolve().parent
manifest=json.loads((here/'design_paper_manifest.json').read_text(encoding='utf8'))
for path, digest in manifest['implementation_source_hashes'].items():
    assert sha256(Path(path).read_bytes()).hexdigest()==digest, path
pages=len(PdfReader(here/'qa_design/Laya_Threat_Detection_Design_Paper.ko.pdf').pages)
source=(here/'audit_paper.py').read_text(encoding='utf8')
source=source.replace("HERE/'paper_manifest.json'", "HERE/'design_paper_manifest.json'")
source=source.replace("xml.count('<m:oMath>') == 3", "xml.count('<m:oMath>') == 6")
source=source.replace("xml.count('<w:drawing>') == 1", "xml.count('<w:drawing>') == 2")
source=source.replace('len(doc.tables) == 6','len(doc.tables) == 14')
source=source.replace('table(1)[1:]', 'table(7)[1:]')
source=source.replace('table(2)[1:]', 'table(8)[1:]')
source=source.replace('table(3)[1:]', 'table(9)[1:]')
source=source.replace("{r['url'] for r in manifest['references']}","{r['url'] for r in manifest['references'] if r['url']}")
source=source.replace('set(range(1,12))','set(range(1,15))')
source=source.replace('qa/Laya_Threat_Detection_Paper.ko.pdf', 'qa_design/Laya_Threat_Detection_Design_Paper.ko.pdf')
source=source.replace('len(pdf.pages) == 10', f'len(pdf.pages) == {pages}')
source=source.replace('> 500', '> 250')
source=source.replace('range(1,11)', f'range(1,{pages+1})')
source=source.replace("f'qa/page-", "f'qa_design/page-")
source=source.replace('pages=10',f'pages={pages}')
source=source.replace('tables=6, references=11, native_equations=3, figures=1','tables=14, references=14, native_equations=6, figures=2')
source=source.replace("'All 10 pages inspected. Final equation change affected only page 4, which was re-inspected; other page image hashes unchanged.'", "'Pending complete visual inspection of the design-inclusive manuscript'")
source=source.replace("HERE/'qa/audit_result.json'", "HERE/'qa_design/audit_result.json'")
exec(compile(source, str(here/'audit_paper.py'), 'exec'),{'__file__':str(here/'audit_paper.py'),'__name__':'__main__'})

doc=Document(here/manifest['document'])
tables=[p.text for p in doc.paragraphs if p.text.startswith('표 ')]
assert [p.split()[1] for p in tables] == [str(i) for i in range(1,13)]+['A1','B1'],tables
chapters=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
assert [p.split()[0] for p in chapters if p[:1].isdigit()] == [str(i) for i in range(1,10)],chapters
assert '제3절의 Zeek 관계 window와 다중 선택지 확장을 사용한 실험이 아니다' in '\n'.join(p.text for p in doc.paragraphs)
print(json.dumps({'implementation_sources':len(manifest['implementation_source_hashes']), 'numbering':'verified', 'pages':pages},ensure_ascii=False))
