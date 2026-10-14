"""Build a Korean research manuscript from frozen experimental evidence.

This script only reads completed results and writes document artifacts.
It does not load models, train, or run inference.
"""
from pathlib import Path
import hashlib, json, math, statistics, zipfile, re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
HERE.mkdir(exist_ok=True)
FIG=HERE/'figures'; FIG.mkdir(exist_ok=True)
STOP=json.loads((ROOT/'stopped/summary.json').read_text(encoding='utf8'))
SPEED=json.loads((ROOT/'speed/verified_summary.json').read_text(encoding='utf8'))
GRID=json.loads((ROOT/'runs/grid_comparison.json').read_text(encoding='utf8'))
assert STOP['laya_completed_repeats']==2 and SPEED['status']=='complete'
assert SPEED['training_resumed'] is False and SPEED['latency_runs']==10

NAMES={
    'encoder_dense':'Encoder Dense', 'encoder_projection':'Encoder Projection',
    'decoder_last':'Decoder Last', 'decoder_flatten':'Decoder Flatten',
    'gpt_lookup':'GPT Lookup', 'gpt_dense':'GPT Dense',
    'bert_none':'BERT None', 'bert_projection':'BERT Projection', 'laya':'Laya'}

REFERENCES=[
 ('L. Daly Manocchio, S. Layeghy, W. W. Lo, G. K. Kulatilleke, M. Sarhan, and M. Portmann. FlowTransformer: A Transformer Framework for Flow-based Network Intrusion Detection Systems. arXiv:2304.14746v1, 2023.', 'https://arxiv.org/abs/2304.14746v1'),
 ('B. Warner et al. Smarter, Better, Faster, Longer: A Modern Bidirectional Encoder for Fast, Memory Efficient, and Long Context Finetuning and Inference. ACL 2025, pp. 2526–2547. DOI 10.18653/v1/2025.acl-long.127.', 'https://aclanthology.org/2025.acl-long.127/'),
 ('Convai Innovations. Laya model card and open model. 본 실험의 가중치 리비전 7b928d828b7b0e022f929d9bd2e44165aa270148. 모델 계열의 최신 설명과 사용 체크포인트를 구분함. 열람 2026-10-07.', 'https://huggingface.co/convaiinnovations/laya'),
 ('D. Almeida. Introducing System One Models & Jev. TypeSafe AI, 2026-09-15. 개발사 기술 소개 자료.', 'https://typesafe.ai/blog/introducing-system-one-models-and-jev'),
 ('T. Deußer, L. Sparrenberg, and R. Sifa. Evaluating and Benchmarking the System One Model Jev. arXiv:2609.37647v1, 2026-09-29. 사전공개 논문.', 'https://arxiv.org/abs/2609.37647v1'),
 ('J. A. dos Santos Barbosa. Calibrated Decision Models for Autonomous Penetration-Testing Harnesses: JEV and Laya as System One Decision Layers for LLM-Driven Pentest Agents. arXiv:2609.28940v1, 2026-09-24. 사전공개 논문.', 'https://arxiv.org/abs/2609.28940v1'),
 ('M. DiGiuseppe and S. Denney. JEV versus LLMs: Accuracy, Cost and Calibration on Seven Political Science Replications. arXiv:2610.06625, 2026-10-05. 사전공개 논문.', 'https://arxiv.org/abs/2610.06625'),
 ('University of Queensland. ML-Based NIDS Datasets. NF-UNSW-NB15-v2 데이터 안내 및 공식 배포. 열람 2026-10-07.', 'https://staff.itee.uq.edu.au/marius/NIDS_datasets/'),
 ('Zeek Project. conn.log. Book of Zeek 9.0.0. 네트워크 연결 로그의 시각, 주체, 대상 및 연결 식별자 정의. 열람 2026-10-07.', 'https://docs.zeek.org/en/current/reference/logs/conn.html'),
 ('MITRE. Working with ATT&CK. 공격 기법 및 완화 지식의 활용 자료. 열람 2026-10-07.', 'https://attack.mitre.org/resources/working-with-attack/'),
 ('L. Daly Manocchio. FlowTransformer source code. 본 실험 고정 커밋 52236c8c9feccdeb44acf578789db7b8b36bacea.', 'https://github.com/liamdm/FlowTransformer/tree/52236c8c9feccdeb44acf578789db7b8b36bacea'),
]

def sha(p):
    with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()

def architecture_figure():
    im=Image.new('RGB',(1800,850),'white'); d=ImageDraw.Draw(im)
    font=lambda n:ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',n)
    def box(x,y,w,h,lines,measured=False):
        d.rounded_rectangle((x,y,x+w,y+h),radius=12,fill='#EAF1F7' if measured else '#F5F5F5',outline='#24465D' if measured else '#666666',width=3)
        for i,line in enumerate(lines):
            b=d.textbbox((0,0),line,font=font(27)); tw=b[2]-b[0]
            d.text((x+(w-tw)/2,y+18+i*42),line,fill='black',font=font(27))
    def arrow(a,b,dashed=False):
        if dashed:
            n=max(1,int(math.dist(a,b)/18))
            for i in range(0,n,2):
                p=(a[0]+(b[0]-a[0])*i/n,a[1]+(b[1]-a[1])*i/n)
                q=(a[0]+(b[0]-a[0])*min(i+1,n)/n,a[1]+(b[1]-a[1])*min(i+1,n)/n)
                d.line((p,q),fill='#555555',width=4)
        else:d.line((a,b),fill='#24465D',width=4)
        angle=math.atan2(b[1]-a[1],b[0]-a[0]); size=16
        pts=[b,(b[0]-size*math.cos(angle-.5),b[1]-size*math.sin(angle-.5)),(b[0]-size*math.cos(angle+.5),b[1]-size*math.sin(angle+.5))]
        d.polygon(pts,fill='#24465D')
    d.text((50,24),'빠른 판정 경로와 비동기 설명 경로',font=font(36),fill='black')
    box(40,130,300,160,['Zeek 로그','conn · dns · tls','수집과 시간 정렬'])
    box(415,130,360,160,['관계 정보를 보존한 window','주체 · 대상 · 시각','필드명과 관측값'])
    box(855,130,385,160,['Laya 판정 모델','언어 인코더와 판정 헤드','위협 점수 및 대응 후보'],True)
    box(1320,130,440,160,['정책 검증과 장비 적용','QoS · 경로 조정','조건부 차단 · 호스트 격리'])
    for a,b in [((340,210),(415,210)),((775,210),(855,210)),((1240,210),(1320,210))]:arrow(a,b)
    box(855,455,385,150,['별도 생성형 LLM','증거 기반 사후 설명','불확실 사례의 추가 분석'])
    box(1320,455,440,150,['운영자 보고 및 감사 기록','원본 증거와 결정 기록','정책 버전 및 실행 결과'])
    arrow((1047,290),(1047,455),True); arrow((1240,530),(1320,530),True)
    box(40,455,735,150,['본 실험에서 검증한 범위','전처리된 8개 흐름 → Laya 이진 판정','자동 대응과 시간·호스트 연관 효과는 후속 검증'],True)
    d.text((50,720),'실선은 판정 경로를, 점선은 비동기 설명 경로를 나타낸다.',font=font(26),fill='#333333')
    d.text((50,765),'자동 대응을 포함한 전체 흐름은 제안 설계이며 실운영 구현 결과가 아니다.',font=font(26),fill='#333333')
    im.save(FIG/'architecture.png')

architecture_figure()
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.9)
sec.left_margin=Cm(2.15); sec.right_margin=Cm(2.15)
sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)

def style_font(style,size,bold=False):
    style.font.name='Malgun Gothic'; style.font.size=Pt(size); style.font.bold=bold
    style.font.italic=False
    style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for k in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+k),'Malgun Gothic')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ko-KR');lang.set(qn('w:eastAsia'),'ko-KR');rpr.append(lang)

for name,size,bold in [('Normal',10,False),('Title',20,True),('Subtitle',11,False),('Heading 1',14,True),('Heading 2',11,True),('Caption',9,False)]:
    style_font(doc.styles[name],size,bold)
normal=doc.styles['Normal'].paragraph_format
normal.space_after=Pt(6);normal.line_spacing=Pt(15);normal.widow_control=True
for node in list(doc.styles.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
for node in list(sec._sectPr.iter(qn('w:docGrid'))):node.getparent().remove(node)
snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0')
doc.styles['Normal'].element.get_or_add_pPr().append(snap)
for name in ['Heading 1','Heading 2']:
    p=doc.styles[name].paragraph_format;p.space_before=Pt(12);p.space_after=Pt(5);p.keep_with_next=True
doc.styles['Title'].paragraph_format.space_after=Pt(9)
doc.styles['Subtitle'].paragraph_format.space_after=Pt(10)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text,style=None):
    q=doc.add_paragraph(text,style);q.paragraph_format.widow_control=True;return q
def h(text,level=1):return doc.add_heading(text,level)
def note(text):
    q=p(text)
    for r in q.runs:r.font.size=Pt(8.5)
    q.paragraph_format.space_after=Pt(8);return q
def formula(text):
    q=p('');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    om=OxmlElement('m:oMath')
    def run(parent,text):
        mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=text;mr.append(mt);parent.append(mr)
    pos=0
    for match in re.finditer(r'([A-Za-zθ∑]+)_([A-Za-zθ]+)',text):
        if match.start()>pos:run(om,text[pos:match.start()])
        sub=OxmlElement('m:sSub');base=OxmlElement('m:e');run(base,match[1]);sub.append(base)
        low=OxmlElement('m:sub');run(low,match[2]);sub.append(low);om.append(sub);pos=match.end()
    if pos<len(text):run(om,text[pos:])
    q._p.append(om)
    return q
def table(caption,headers,rows,widths):
    cap=p(caption,'Caption');cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    def fillrow(row,vals,ishead=False,index=0):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ishead:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,(cell,val,w) in enumerate(zip(row.cells,vals,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(val);q=cell.paragraphs[0]
            q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after=Pt(2);q.paragraph_format.space_before=Pt(2);q.paragraph_format.line_spacing=1.12
            q.paragraph_format.keep_with_next=ishead
            for r in q.runs:r.font.size=Pt(8.5);r.font.bold=ishead
            pr=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6E6E6' if ishead else ('F7F7F7' if index%2 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side,n in [('top',75),('bottom',75),('left',85),('right',85)]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),str(n));el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
    fillrow(t.rows[0],headers,True)
    for i,row in enumerate(rows):fillrow(t.add_row(),row,index=i)
    return t
def link(q,text,url):
    rel=q.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'24465D');pr.append(col)
    r.append(pr);el=OxmlElement('w:t');el.text=text;r.append(el);hl.append(r);q._p.append(hl)

p('Laya 아키텍처 기반 저지연 언어모델 위협탐지와 자동 대응 설계','Title')
p('Low Latency Language Model Threat Detection and Response with the Laya Architecture','Subtitle')
note('공개 네트워크 흐름 데이터 기반 탐색적 실증 연구  |  2026년 10월 7일')
h('초록')
p('네트워크 위협의 초기 대응은 여러 로그의 문맥을 해석하는 능력과 짧은 판정 지연을 동시에 요구한다. 생성형 대규모 언어모델을 이용해 위협 설명까지 동기적으로 생성하면 대응 결정이 설명 완료를 기다리는 구조가 된다. 본 연구는 언어로 기술된 상태와 판정 기준을 입력받되 텍스트 생성 없이 선택지 점수를 반환하는 Laya를 빠른 위협 판정 경로에 배치하고, 별도 생성형 언어모델이 증거 설명을 비동기적으로 담당하는 구조를 제안한다. 핵심 관점은 언어 사전학습을 활용하는 판정 인터페이스를 유지하면서 응답 경로에서 순차적인 설명 생성을 제거하는 데 있다.')
p('실증 평가에서는 NF-UNSW-NB15-v2의 2,390,275개 흐름을 사용해 FlowTransformer의 여덟 비교 설정과 Laya를 평가했다. Laya의 완료된 두 학습 반복은 평균 F1 93.90%, 공격 탐지율 99.948%, 오탐률 0.946%를 보였다. 같은 RTX 4090에서 입력 준비를 포함한 단일 요청 지연 p95는 80.03ms와 76.37ms로, 사전에 정한 추론 경로의 200ms 목표를 충족했다. 그러나 최상 비교군 Encoder Projection은 F1 94.49%, 오탐률 0.832%, p95 16.22ms로 더 나은 F1·오탐·지연을 보였다. 또한 실제 호스트·시간 관계에 의존하는 연관 탐지, 미학습 공격, 대응 방책 선택 및 차단 효과는 평가하지 않았다. 결과는 Laya가 저지연 언어 기반 판정 계층의 실현 가능성을 갖는다는 근거를 제공하며, 언어모델의 장점을 실시간 대응에 연결하려는 설계적 진전의 방향을 지지한다. 다만 언어 이해 능력의 보존과 기존 탐지기를 능가하는 종합 성능은 별도의 검증 과제다.')
note('주제어  Laya  비생성형 의사결정 모델  네트워크 위협탐지  로그 연관 분석  저지연 추론  자동 대응')

h('1 서론')
p('보안 운영에서 단일 연결의 패킷 수나 전송량만으로는 정상과 공격을 구분하기 어려운 경우가 있다. 같은 주체의 DNS 조회, 반복 연결, 대상 변경, 전송량 증가를 시간과 관계에 따라 결합해야 위협 가능성을 판단할 수 있다. 언어모델은 서로 다른 로그 형식과 자연어로 기술된 정책을 한 입력 안에서 다룰 수 있어 이러한 연관 분석의 후보가 된다. 그러나 설명문 생성, 다단계 추론, 원격 호출을 모두 초기 차단 이전에 수행하는 구성에서는 각 단계의 지연이 대응 경로에 누적된다. 이 문제는 모든 LLM이 항상 느리다는 일반 명제가 아니라, 설명 생성까지 기다리는 시스템 구성의 문제로 정의한다.')
p('Jev와 Laya는 상태와 질문을 받아 정해진 선택지 또는 점수의 분포를 반환하는 접근을 제시한다[3–5]. 이 접근은 자연어 입력과 유연한 판정 기준을 활용하면서 출력 생성을 제한하므로, 프로그램이 빠른 결정을 직접 소비할 수 있는 가능성을 제공한다. 본 연구는 이 방향을 네트워크 방어에 적용한다. 빠른 모델은 위협 여부와 대응 후보를 판정하고, 설명이 필요한 경우 별도 LLM이 원본 증거와 판정 기록을 토대로 상세 분석을 수행한다. 초기 대응이 완료된 뒤 설명을 제공할 수 있으므로, 설명의 길이와 초기 판정 지연을 분리할 수 있다.')
p('본 연구의 기여는 공개 흐름 데이터에서 Laya를 학습한 재실행 결과, 같은 장치에서 측정한 입력 준비 포함 지연, 그리고 판정·정책 적용·설명을 분리한 방어 구조의 제안이다. 새로운 Transformer 연산이나 최초의 System One 보안 통합을 주장하지 않는다. 이미 유사한 보안 의사결정 계층이 논의됐으므로[6], 기여 범위는 방어용 흐름 분류에서의 확인 가능한 실측과 연관 탐지 검증 조건의 정리에 둔다.')

h('2 관련 연구와 연구 질문')
h('2 1 흐름 Transformer와 언어 사전학습',2)
p('FlowTransformer는 입력 인코딩, Transformer, 분류 헤드를 교체하며 네트워크 흐름 분류를 평가하는 프레임워크다[1,11]. 여러 흐름의 표현을 함께 처리하므로 수치 패턴의 관계를 학습할 수 있지만, 모델에 여러 행을 넣었다는 사실만으로 시간·호스트·공격 단계의 의미적 연관을 활용했다고 볼 수는 없다. 본 연구에서는 arXiv v1의 Table III UNSW-NB15 설정을 기준으로 네 계열의 각 두 설정을 재실행했다. 논문에 등장하는 GPT·BERT는 저자 코드의 구조를 처음부터 학습한 비교군이며, 상용 GPT 서비스나 사전학습 Qwen과 다르다.')
p('ModernBERT는 분류와 검색을 위한 양방향 언어 인코더로, 효율적인 추론과 긴 문맥 처리를 고려한 구조를 제시한다[2]. Laya는 이러한 언어 사전학습 인코더와 판정 헤드를 결합한다[3]. 따라서 Laya와 작은 Encoder의 차이는 Encoder 사용 여부보다 사전학습, 입력 표현, 모델 용량 및 학습 목적에 있다. 언어 사전학습 가중치를 사용한다는 사실은 장점의 활용 경로를 제공하지만, 보안 파인튜닝 뒤 의미 이해나 미학습 공격 일반화가 유지됐다는 실증 결과를 대신하지 않는다.')
h('2 2 Jev와 Laya의 의사결정 인터페이스',2)
p('TypeSafe의 Jev 소개 자료는 비생성형 구조화 결정과 낮은 응답 지연을 지향한다[4]. 개발사의 특정 업무 벤치마크에서 보고한 가속 배율은 본 연구의 네트워크 탐지 결과와 직접 결합하지 않는다. Jev와 Laya는 유사한 인터페이스 방향을 공유하지만, Jev의 내부 구조가 Laya와 동일하다고 가정하지 않는다. 또한 형식에 맞는 출력을 반환하는 성질과 판정 내용이 올바른 성질은 구분해야 한다.')
p('Deußer 등의 사전공개 연구는 Jev를 37개 일반 언어 과제에서 평가했다[5]. 일부 분류·추론 과제에서 높은 성능을 보고했지만, 이진 판정의 임계값과 과제별 보정 문제도 지적했다. 해당 연구의 평균 클라이언트 지연은 동시 요청 32개 조건에서 0.36초이며 네트워크 시간이 포함됐다. 이 값은 본 연구의 로컬 Laya p95와 측정 경계가 달라 속도 순위 산정에 사용하지 않는다. DiGiuseppe와 Denney의 사회과학 과제 연구도 정확도·보정·비용의 이점이 과제와 비교 조건에 따라 달라짐을 보고했다[7]. 이들 연구는 의사결정 모델의 언어 활용 가능성을 뒷받침하지만 네트워크 연관 탐지 능력을 입증한 자료는 아니다.')
p('Barbosa는 Jev·Laya를 LLM 기반 침투시험의 의사결정 계층에 배치하는 구성을 논의했다[6]. Jev 유무를 각각 한 번 실행한 사례 연구이며 Laya의 보안 성능을 직접 비교하지 않았다는 한계를 명시했다. 본 연구는 이 흐름을 방어용 네트워크 탐지로 확장하되, 아키텍처의 타당성과 실험으로 확인한 성능을 분리한다.')
h('2 3 검증할 주장과 현재 근거',2)
table('표 1 연구 질문과 본 연구의 증거 범위', ['연구 질문','현재 상태','판단 근거'],[
 ['Laya 판정 경로가 200ms 이내인가','측정 완료','두 체크포인트 p95 76.37–80.03ms'],
 ['기존 흐름 탐지기보다 우수한가','종합 우위 미입증','F1·오탐·지연 불리, 탐지율 소폭 증가'],
 ['언어 이해와 연관 탐지 장점이 유지되는가','후속 검증 필요','문맥 대조·언어 능력 평가 없음'],
 ['생성형 LLM보다 빠르게 같은 위협을 찾는가','직접 비교 없음','동일 입력의 생성형 LLM 측정 없음'],
 ['자동 대응이 공격 확산을 줄이는가','제안 단계','정책 실행 및 피해 억제 실험 없음'],
 ],[6.1,3.0,7.5])

h('3 Laya 기반 위협탐지와 자동 대응 구조')
h('3 1 실험에 사용한 Laya 판정 모델',2)
p('영어 Laya 체크포인트는 ModernBERT-large 기반의 양방향 인코더와 두 층의 판정용 Transformer, 선택지 마커 점수기 및 별도 act/escalate 헤드를 포함한다[3]. 사용한 모델의 전체 파라미터 수는 421,293,827개다. 상태, 질문, 선택지 설명을 토큰열로 구성하고 각 선택지의 마커 위치에 대응하는 점수에 softmax를 적용한다. 본 실험은 원래의 false/true 질문 형식을 유지하고 마지막 흐름이 공격인지 묻는 이진 판정으로 제한했다. 설명문이나 공격 이름을 생성하지 않았다.')
p('입력은 여덟 흐름의 전처리된 특성, 열 이름, 마지막 행을 판정하라는 영어 지시문이다. 수치값은 float32가 보존되도록 유효숫자 아홉 자리의 CSV로 직렬화했다. 모델 입력 상한은 8,192토큰으로 설정했고, 완료된 두 학습·평가 실행에서 관측한 최대 길이는 2,080토큰과 2,074토큰이었다. 이 상한 설정은 모든 문맥 길이에서 성능이 유지된다는 보장이 아니며, 실제 관측 길이를 넘어선 성능은 평가하지 않았다.')
formula('p(attack | W, q) = exp(s_true) / (exp(s_false) + exp(s_true))')
note('W는 입력 window, q는 판정 질문이며 s는 선택지 점수다. 본 실험의 결정은 p > 0.5이고 이 점수를 실제 환경에서 보정된 위협 확률로 간주하지 않는다.')
h('3 2 빠른 판정과 비동기 설명의 분리',2)
p('제안 구조에서는 Zeek의 연결·DNS·TLS 로그를 시간과 식별자에 따라 연결한다. Zeek 연결 로그의 ts, uid, 주체·대상 주소 등은 관계 구성의 근거가 될 수 있다[9]. 실운영 입력에는 동일 주체를 일관되게 연결할 수 있는 익명 식별자, 사건 간 시간차, 서비스와 대상의 관계를 남긴다. window는 시간 범위와 최대 사건 수를 함께 제한하며, 결정 시점에 관측된 정보만 사용한다. 전체 연결 종료 후에만 얻을 수 있는 값을 조기 대응 시점의 입력처럼 제공해서는 안 된다.')
doc.add_picture(str(FIG/'architecture.png'),width=Cm(16.55))
doc.paragraphs[-1].paragraph_format.line_spacing=1.0
doc.paragraphs[-1].paragraph_format.keep_with_next=True
p('그림 1 빠른 위협 판정과 비동기 설명을 분리한 제안 구조','Caption')
p('빠른 경로는 입력 window를 Laya에 전달해 위협 점수와 허용된 대응 후보를 산출하고, 정책 검증을 거쳐 장비에 적용한다. 느린 경로는 동일한 원본 증거, 모델 버전, 판정 점수와 실제 실행 결과를 별도 생성형 LLM에 전달한다. 이 경로는 공격 가능성, 근거 로그 및 대응 이유를 설명하되, 이미 실행된 결정을 사후 정당화하는 설명을 만들지 않도록 관측 사실과 추론을 구분한다. 낮은 확신도나 정책 충돌은 추가 분석 대상으로 남긴다.')
formula('T_total = T_obs + T_prep + T_queue + T_infer + T_policy + T_apply')
p('위 식에서 관측 지연은 필요한 사건이 생성·수집되는 시간, 입력 준비 지연은 정규화와 직렬화 시간, 추론 지연은 모델 계산 시간이다. 비동기 설명 시간은 빠른 경로의 합에 포함하지 않는다. 이번 속도 측정은 이미 전처리된 window의 입력 조립·직렬화·토큰화부터 CPU 판정값 반환까지다. 공격 발생부터 장비 차단까지의 전체 지연은 측정하지 않았다.')
h('3 3 대응 방책과 정책 제약',2)
p('실험 모델의 실제 출력은 위협 여부 하나다. 대응 방책 선택은 후속 확장으로 제안한다. 선택지는 관측 유지, QoS 제한, 경로 변경, 특정 연결 또는 패킷 조건 차단, 호스트 연결 제한, 추가 분석 요청으로 구성할 수 있다. 모델이 내놓는 값은 운영 코드가 정의한 기능 식별자이며, 임의의 장비 명령 문자열을 생성하는 방식은 사용하지 않는다. 이 선택 공간을 둔다고 대응 효과가 자동으로 학습되는 것은 아니므로, 상태별 허용 조치와 조치 결과를 포함한 별도 정답 자료가 필요하다.')
p('조치는 위협 점수 하나만으로 결정하지 않고 정상 서비스 손실, 공격 피해, 자산 중요도, 조치 범위 및 복구 가능성을 함께 반영한다. QoS나 한 연결의 제한과 호스트 전체 격리는 비용이 다르므로 별도 기준으로 검증한다. 정책 계층에는 대상 일치 검사, 허용 기능 목록, 적용 기간, 중복 실행 방지 및 해제 절차를 둔다. 이는 본 논문의 제안 설계이며 실제 라우터에 적용한 성능 결과가 아니다.')
p('ATT&CK와 mitigation 자료는 공격 기법과 대응 개념을 연결하는 지식으로 활용한다[10]. 이 지식만으로 개별 로그의 정상·공격 정답이나 최적 QoS 값을 만들 수는 없다. 공개 데이터셋이나 관측 시나리오의 라벨을 유지하고, 지식 매핑은 근거와 함께 별도로 기록해야 한다. 판정 점수의 확률 보정과 조치별 오류 비용도 독립 검증 구간에서 정해야 한다.')

h('4 공개 데이터 기반 학습과 평가 방법')
h('4 1 데이터와 분할',2)
p('실험은 UQ 공식 NF-UNSW-NB15-v2의 2,390,275개 흐름을 사용했다[8]. 원본 행 순서에서 앞 2,151,248행을 학습 구간으로, 뒤 239,027행을 평가 구간으로 나눴다. 평가 구간에는 공격 16,271개와 정상 222,756개가 있다. 저자 코드가 선택한 37개 특성을 사용하고 IP 주소, 정답 Label과 공격 이름 Attack은 입력에서 제외했다. 수치 변환 및 빈도가 낮은 범주 처리는 학습 구간에서 적합했다. 입력은 연속 여덟 행이며 마지막 행의 이진 라벨이 정답이다.')
p('원문은 90:10 분할을 기술하지만 공개 UNSW 예제에는 97.5:2.5가 사용돼 있다[1,11]. 본 연구는 90:10을 채택하고 구체적인 행 경계와 해시를 남겼다. 원문의 정확한 평가 행과 최적 학습률·seed가 공개되지 않았으므로 본 결과를 논문 발표값과 완전히 동일한 재현이라고 부르지 않는다. 또한 v2 CSV에는 흐름 시작 시각이 없으며 행의 인접성이 같은 호스트나 공격 단계의 연결을 보장하지 않는다. 이 평가 단위는 시간 기반 Zeek 연관 window와 구별된다.')
h('4 2 학습 설정과 완료 범위',2)
p('비교군은 Shallow Encoder, Shallow Decoder, GPT 구조, BERT 구조의 각 두 설정이다. seed 20261005, 20261006, 20261007별 표본 스케줄을 사전에 고정해 모든 모델에 동일하게 공급했다. 유효 배치는 128개로 정상·공격 각 64개이며, epoch당 64개 배치를 추출한다. 따라서 한 epoch는 학습 구간 전체를 한 번 순회한다는 뜻이 아니라 8,192회 표본 추출이다. 최대 20 epoch와 원본 코드의 최소 개별 배치 손실에 기반한 조기 종료를 사용했다.')
p('Laya는 학습률 0.00001의 Adam과 false/true 점수에 대한 교차엔트로피로 인코더 및 판정 경로를 파인튜닝했다. 이진 판정에 사용하지 않는 act/escalate 헤드는 고정했으며 학습 가능한 파라미터는 421,029,889개다. micro-batch 8과 누적 16회로 유효 배치 128을 맞췄다. 가중치는 float32, 계산은 BF16 autocast와 SDPA를 사용했다. 기본 모델의 RLCD 학습 설명과 이번 지도학습을 혼동하지 않으며, 파인튜닝 뒤 확률 보정은 수행하지 않았다.')
formula('L(θ) = −(1/N) ∑_i log p_θ(y_i | W_i, q)')
p('비교군은 학습률 0.01, 0.001, 0.0005, 0.0001, 0.00001의 다섯 값에서 각 세 번 실행했다. 기본 BERT 설정의 수렴 실패 후 추가 탐색 조건을 정했으며 Laya의 첫 테스트 평가 전에 고정했다. 비교군 총 120회 중 119회가 완료되고 한 번은 수치 발산했다. Laya는 세 번 중 두 번이 각각 17·12 epoch에서 종료된 뒤 전체 평가를 마쳤다. 세 번째는 학습 중단으로 평가하지 않았으며 본문에 성적을 부여하지 않았다. 속도 보강 시에는 저장된 가중치의 추론만 수행했다.')
p('각 비교군의 학습률은 세 반복 평균 테스트 F1로 선택했다. 이 선택에는 테스트 기반 낙관성이 있고 Laya는 한 학습률만 사용했으므로 탐색량도 비대칭이다. 결과는 관측된 비교이며 통계적으로 확정된 모델 일반의 우열로 해석하지 않는다. 테스트 성적이 가장 좋은 epoch를 고른 것은 아니며, 주 평가 체크포인트는 정해진 학습 종료 시점의 가중치다.')
h('4 3 정확도와 속도 지표',2)
p('탐지 지표는 F1, 공격 탐지율, 정상 흐름의 오탐률과 혼동행렬이다. 평균과 표본 표준편차를 보고하고, 반복 수가 적어 표준편차를 신뢰구간으로 해석하지 않는다. 모든 공격 종류가 학습 구간에 포함돼 있으므로 이 결과는 미학습 공격 평가가 아니다. 학습 구간과 같은 특성값이 있는 평가 행을 제외한 하위 집단도 별도로 집계하되, 이를 새로운 독립 테스트셋이나 중복 제거 재학습 결과로 부르지 않는다.')
p('추론은 같은 RTX 4090에서 모델을 순차 실행했다. 비교군은 이미 선택한 학습률의 첫 seed, Laya는 첫 완료 seed를 사용했으며 두 번째 Laya는 단일 요청 지연만 추가 확인했다. 각 배치 크기 1·8·128에 네 배치를 워밍업한 뒤 고정된 50개 무작위 학습 배치를 각각 네 번 측정했다. 처리량은 배치별 네 반복의 중앙값을 50개 배치에 대해 평균한 시간으로 계산했다. p50·p95는 입력 준비를 포함한 단일 요청 200개의 원시 지연에서 직접 계산했다.')
p('단일 요청은 같은 50개 window의 네 반복이며 독립적인 200개 입력이 아니다. 모델 로딩, 최초 컴파일, 원시 로그 파싱·수치 정규화, window 수집 대기, 동시 요청 큐, 네트워크 및 장비 적용은 제외했다. 비교군은 TensorFlow 2.15.1 float32, Laya는 PyTorch 2.8.0+cu128의 BF16 autocast를 사용했다. WSL2에서 화면 출력에도 쓰는 GPU를 사용했고 측정 직전 네 차례 사용률은 0~2%였다. 운영체제 간섭을 완전히 제거한 전용 장비의 장시간 부하 시험은 아니다.')

h('5 실험 결과')
h('5 1 탐지 성능',2)
acc=[]
for name,model in GRID['models'].items():
    b=model['best_mean'];acc.append([NAMES[name],model['best_mean_lr'],3,f'{b["f1_mean"]*100:.2f} ± {b["f1_sd"]*100:.2f}',f'{b["recall_mean"]*100:.3f}',f'{b["fpr_mean"]*100:.3f}'])
l=STOP['laya_aggregate'] if 'laya_aggregate' in STOP else STOP.get('laya_summary')
# Aggregate directly from preserved per-seed records to avoid key-schema assumptions.
lr=STOP['laya_rows']; fm=statistics.mean(r['f1'] for r in lr);fs=statistics.stdev(r['f1'] for r in lr)
acc.append(['Laya','0.00001',2,f'{fm*100:.2f} ± {fs*100:.2f}',f'{statistics.mean(r["recall"] for r in lr)*100:.3f}',f'{statistics.mean(r["fpr"] for r in lr)*100:.3f}'])
table('표 2 선택한 학습률에서의 탐지 성능', ['모델','학습률','반복','F1 %','탐지율 %','오탐률 %'],acc,[4.05,2.0,1.0,3.55,3.0,3.0])
note('F1은 평균 ± 표본 표준편차다. 비교군은 테스트 F1로 학습률을 선택했으며 Laya와 탐색량이 다르다. GPT·BERT는 사전학습 상용 언어모델이 아닌 저자 코드의 비교 구조다.')
p('Laya의 평균 F1은 93.90%로 최상 비교군 Encoder Projection의 94.49%보다 0.59%p 낮았다. 공격 탐지율은 99.948%로 비교군 99.756%보다 높았지만 오탐률은 0.946%로 비교군 0.832%보다 높았다. 같은 두 seed만 비교해도 F1은 93.90% 대 94.44%로 Laya가 0.55%p 낮았다. 평균 누락은 46.0개에서 8.5개로 감소한 반면 평균 오탐은 1,863.5개에서 2,107.0개로 증가했다. 놓치는 공격을 줄이는 방향의 변화와 정상 서비스에 부담을 주는 변화가 함께 관측됐다.')
table('표 3 Laya와 Encoder Projection의 공통 seed 혼동행렬', ['모델과 seed','TP','FP','TN','FN'],[
 ['Laya 20261005',16267,1907,220849,4],['Encoder 20261005',16239,1814,220942,32],
 ['Laya 20261006',16258,2307,220449,13],['Encoder 20261006',16211,1913,220843,60],
 ],[5.4,2.8,2.8,2.8,2.8])
p('Laya의 평균 정밀도는 88.54%이고 정상 흐름 10,000개당 평균 오탐은 94.6개다. 이를 실제 호스트 차단 수로 환산할 수는 없지만, 공격 점수만으로 강한 차단을 자동 적용하기에는 조치 단위의 검증이 필요함을 보여준다. BERT 계열의 큰 표준편차와 높은 오탐률은 일부 반복이 거의 모든 흐름을 공격으로 판정한 결과다. 이 수렴 실패를 BERT 계열 전반의 한계로 일반화하지 않는다.')
p('평가 행 59,793개인 약 25.02%는 학습 구간에도 동일한 37개 특성 조합이 존재했다. 이들을 제외한 179,234개에서 Laya 평균 F1은 92.21%, 탐지율은 99.952%, 오탐률은 1.250%였다. 동일 특성의 존재는 같은 원본 패킷이나 연결의 중복 수집을 확인한 것은 아니다. 하위 집단의 공격 비율도 달라지므로 F1 차이를 모두 암기 효과로 설명하지 않는다.')
h('5 2 추론 지연과 배치 처리량',2)
speedrows=[]
for name in NAMES:
    d=SPEED['timings'][f'{name}/20261005'];o=d['batch1_serialization_inclusive']
    speedrows.append([NAMES[name],f'{o["p50_ms"]:.2f}',f'{o["p95_ms"]:.2f}',f'{d["batch8"]["flows_per_second"]:.1f}',f'{d["batch128"]["flows_per_second"]:.1f}'])
table('표 4 같은 RTX 4090에서 측정한 추론 지연과 처리량', ['모델','p50 ms','p95 ms','배치 8 판정/초','배치 128 판정/초'],speedrows,[4.2,2.5,2.5,3.7,3.7])
note('p50·p95는 입력 준비를 포함한 단일 요청 지연이다. 처리량은 사전 인코딩한 입력부터 CPU 점수 반환까지이며 직렬화·토큰화는 제외한다. Laya 배치 128은 micro-batch 16을 순차 실행했다. 한 판정은 8개 흐름 window의 마지막 흐름에 대한 판정이다.')
p('Laya 첫 체크포인트의 p50은 29.61ms, p95는 80.03ms였고 두 번째는 각각 28.53ms와 76.37ms였다. 각 체크포인트의 200회 요청은 모두 200ms 이내였으며 관측 최대 지연은 각각 112.89ms와 88.93ms였다. 이 결과는 워밍업된 로컬 추론 경로에서 정한 지연 목표의 충족을 보여준다. 다만 p95의 표본 크기가 작고 장시간 부하·동시 요청 조건이 아니므로 운영 SLA를 보장하지 않는다.')
p('Encoder Projection의 p95는 16.22ms로 Laya 첫 체크포인트보다 약 4.9배 짧았다. 따라서 Laya가 이번의 작은 탐지기보다 빠르다는 주장은 성립하지 않는다. 모델 크기도 약 35만 대 4억 2천만 파라미터로 크게 다르다. 이 차이에 입력 길이와 프레임워크가 함께 작용했으므로 Projection 하나만의 인과 효과를 추정하지 않는다. 생성형 LLM의 같은 질문에 대한 응답 지연은 측정하지 않아, 본 실험으로 LLM 대비 가속 배율을 산출할 수는 없다.')
p('Laya의 사전 인코딩 입력 처리량은 배치 8에서 43.3판정/초, 배치 128에서 44.5판정/초였다. 배치를 크게 해도 물리 배치 16으로 나누어 처리하는 조건에서는 증가 폭이 작았다. 단일 요청이 빠르다는 사실과 많은 요청이 동시에 도착할 때의 대응 능력은 다르다. 운영에서는 입력 준비, 배치 채움 대기, 큐 적체를 포함한 지연과 용량을 함께 측정해야 한다.')

h('6 논의')
h('6 1 진일보한 방향의 의미',2)
p('본 연구가 지지하는 진전은 언어모델 기반 보안 시스템의 역할 배분에 있다. 위협의 설명을 생성하는 능력과 즉시 실행할 결정을 반환하는 능력을 별도 경로로 구성하면, 초기 대응이 설명 생성의 완료를 기다릴 필요가 없다. Laya는 자연어로 기술한 질문과 선택지를 처리하는 구조를 유지하면서 로컬 추론 p95 200ms 이내를 달성했다. 이는 언어모델의 활용 범위를 실시간 판정으로 넓히려는 방향의 실현 가능성을 보여준다.')
p('여기서 언어모델의 장점을 유지한다는 표현은 사전학습 인코더와 언어 기반 판정 인터페이스를 유지한다는 구조적 의미로 한정한다. 기존 LLM과 같은 수준의 보안 추론, 새로운 공격의 이해, 정책 변경에 대한 일반화가 유지됐다고 주장하지 않는다. 원래 Laya의 지시문 처리 능력도 보안 파인튜닝으로 달라질 수 있으므로, 학습 전후의 언어·보안 평가가 필요하다. 설명 생성을 제거한 구조적 이점은 방향의 근거이며, 실제 탐지 우위는 별도 실험의 대상이다.')
p('기존 Encoder Projection은 수치 흐름 분류에서 더 좋은 F1·오탐·지연을 보였다. 따라서 이번 조건에서의 대체 근거는 부족하다. 그러나 이 결과만으로 두 모델의 의미적 연관 탐지 능력까지 비교했다고 볼 수도 없다. Laya가 더 유리할 가능성이 있는 과제는 형식이 다른 사건과 자연어 정책을 함께 해석해야 하는 판정이며, 그 가능성은 과제에 맞는 데이터로 검증해야 한다. 속도 목표 충족, 탐지 성능 개선, 연관 의미 이해의 보존을 서로 다른 주장으로 다루는 것이 필요하다.')
h('6 2 현재 평가가 설명하지 못하는 부분',2)
p('첫째, 입력에서 호스트 식별자를 제외하고 시작 시각을 제공하지 않았으므로 다단계 공격의 주체·시간 관계를 충분히 나타내지 못했다. 둘째, 단일 흐름 대조 학습이나 문맥 제거 평가가 없어 성적이 문맥 활용에서 비롯됐는지 알 수 없다. 셋째, 공격 계열을 학습에서 제외하지 않아 미학습 위협 탐지율을 제시할 수 없다. 이 제약은 두 모델 모두에 적용되며 어느 한 모델의 능력 부재를 뜻하지 않는다.')
p('넷째, Laya는 두 반복만 완료됐고 비교군의 테스트 기반 학습률 선택과 비대칭 탐색이 있다. 다섯째, 데이터의 행 순서 분할이 원문과 정확히 같은지 확인할 수 없고, 동일 특성 중복도 존재한다. 여섯째, 정밀도·optimizer 세부값·학습 GPU가 다르며 BERT의 모델 전체 XLA JIT을 비활성화한 재현 차이가 있다. 속도는 같은 RTX 4090에서 다시 측정했지만 커널과 프레임워크까지 동일하게 통제한 비교는 아니다.')
p('일곱째, Jev의 수치는 관련 문헌의 결과이며 이번 장치와 데이터에서 실행한 결과가 아니다. 기존 생성형 LLM이 같은 네트워크 위협 판정에서 얼마나 느린지에 대한 대조군도 없다. 모든 LLM이 수초 이상 걸린다거나 Jev·Laya가 일반적으로 200ms를 보장한다고 확대하지 않는다. 짧은 선택지 출력이나 next-token 점수 방식의 LLM도 비교해야 한다. 마지막으로, 형식이 제한된 출력도 잘못된 공격 판단이나 잘못된 조치를 만들 수 있으며 확률 보정은 도메인별로 다시 확인해야 한다.')

h('7 연관 탐지와 자동 대응을 위한 후속 검증')
h('7 1 관계를 보존하는 공개 네트워크 평가',2)
p('후속 평가의 기본 단위는 시간과 주체가 유지된 네트워크 사건 묶음으로 정한다. 공개 PCAP 또는 시각·주소가 보존된 로그에서 Zeek 기록을 만들고, 분할 전에 공격 시나리오와 수집 구간을 구분한다. 같은 공격 실행이나 중첩 window가 학습·검증·평가를 가로지르지 않게 하며 주소는 분할 누수를 피하면서 관계를 유지하는 익명 식별자로 표현한다. 단순히 IP를 제거하면 일반화 누수는 줄일 수 있지만, 필요한 관계 정보까지 사라질 수 있다는 점을 함께 고려한다.')
p('문맥이 필요한 평가 집단은 마지막 사건만으로 정답이 결정되지 않는 사례로 구성한다. 예를 들어 비슷한 연결이 정상 관리 작업의 일부인지, 같은 주체의 탐색·반복 통신과 결합된 공격 단계인지 구분하도록 한다. 같은 형식의 정상 업데이트·백업 사례를 포함하고, 문맥을 바꿀 때 정답도 전문가 또는 재현 시나리오에 따라 검증한다. 단순히 문맥을 섞었을 때 모델 점수가 달라지는 현상만으로 올바른 연관 추론을 입증하지 않는다.')
table('표 5 후속 실험에서 필요한 대조 조건', ['검증 목적','대조 조건','핵심 지표'],[
 ['문맥의 실제 기여','단일 사건 학습 대 전체 window 학습','같은 오탐률의 사건 탐지율'],
 ['관계 정보의 기여','호스트·시간 관계 유지 대 통제된 관계 변경','정답 변화에 대한 판정 일관성'],
 ['미학습 공격 일반화','공격 계열·시나리오를 통째로 평가에 보류','미학습 공격 탐지율과 정상 오탐'],
 ['언어 장점의 유지','원본·파인튜닝 모델과 사전학습 제거 대조','정책 바꿔쓰기·새 필드·새 조합 성능'],
 ['생성형 모델 대비 효율','설명 생성·짧은 결정·선택지 점수 방식','같은 정확도 조건의 p95와 처리량'],
 ['자동 대응 효과','탐지만 수행 대 정책 적용을 포함한 재생','확산 범위·차단 시간·정상 서비스 손실'],
 ],[3.6,7.0,6.0])
h('7 2 학습과 보정 및 정책 평가',2)
p('학습은 이진 위협 판정에서 시작하되, 연관 증거가 필요한 사례와 정상 유사 사례를 늘리는 방식으로 확장한다. 위협 유무, 위험 수준, 허용 대응 후보의 목표를 구분하고 대응 후보에는 미조치·추가 분석도 포함한다. ATT&CK 매핑이나 LLM 설명을 자동 정답으로 확정하지 않고, 공개 시나리오의 실제 라벨과 전문가 확인으로 검증한다. 자연어 표현 다양성을 유지하는 학습이나 기존 모델의 출력 분포를 보존하는 증류는 후보 방법이며, 효과를 입증하려면 해당 항목을 제거한 대조 실험이 필요하다.')
p('확률 보정과 조치 임계값은 독립 검증 구간에서 결정하고 테스트셋을 사용하지 않는다. 같은 오탐률에서의 탐지율, PR-AUC, Brier score, ECE, 판정 보류 비율과 위험을 함께 보고한다. 조치 정답의 일치율뿐 아니라 정상 연결 단절, 서비스 지연 증가, 공격 확산 감소도 재생 환경에서 측정해야 한다. 정책이 정한 기능 목록과 모델의 선택지를 같은 버전으로 기록해 변경 후에도 결과를 추적할 수 있게 한다.')
p('지연 평가는 정상 요청률과 공격 중 급증 구간을 나눠 수행한다. 토큰 수, window 길이, 동시 요청 수, 실제 물리 배치, 큐 길이별 p50·p95·p99를 기록하고 공격 발생부터 조치 적용까지를 계측한다. Laya·Jev·생성형 LLM을 비교할 때에는 입력 증거와 판정 목표를 맞추고, 로컬 계산 시간과 네트워크를 포함한 서비스 지연을 별도로 보고한다. 이 조건이 갖춰져야 언어모델의 장점을 유지하면서 대응을 앞당겼다는 주장을 직접 검증할 수 있다.')

h('8 결론')
p('본 연구는 Laya 아키텍처를 이용해 언어 기반 위협 판정을 빠른 대응 경로에 배치하고, 별도 생성형 LLM이 상세 설명을 담당하는 구조를 제안했다. 공개 네트워크 흐름 데이터에서 학습한 두 Laya 체크포인트는 RTX 4090의 입력 준비 포함 단일 요청 p95 76.37–80.03ms를 보여, 이번 추론 경로의 200ms 목표를 충족했다. 결과는 언어 사전학습을 활용하는 비생성형 판정 계층을 저지연 보안 시스템에 넣을 수 있다는 실증 근거다.')
p('동시에 Laya의 평균 F1과 오탐률은 최상 비교군보다 불리했고 추론도 더 느렸다. 따라서 현재 결과로 기존 탐지기보다 우수한 모델을 완성했다고 결론내릴 수 없다. 본 연구가 제시하는 진일보는 설명 생성에 묶여 있던 초기 판정을 분리하고, 언어 기반 정책 해석의 가능성과 짧은 응답 경로를 함께 추구하는 설계 방향이다. 그 방향이 실제 연관 위협 탐지와 공격 확산 억제로 이어지는지는 관계를 보존한 데이터, 미학습 공격 대조, 확률 보정 및 정책 실행 평가로 확인해야 한다.')

h('참고문헌')
for i,(text,url) in enumerate(REFERENCES,1):
    q=p(f'[{i}] {text} ');q.paragraph_format.space_after=Pt(7)
    for r in q.runs:r.font.size=Pt(8.5)
    link(q,'원문',url)

h('부록 A 재현 자료와 수치 추적')
p('본 논문의 탐지 수치는 학습 중단 시점의 저장된 예측에서 산출했고, 속도는 이후 저장된 체크포인트의 추론만으로 보강했다. 이전 v1–v4 실험과 서로 다른 Qwen 평가값을 동일 조건의 결과로 혼합하지 않았다. 두 Laya의 전체 평가 행 번호·정답·예측 확률은 로컬 파일로 보존했으며 SHA-256과 혼동행렬을 검증했다. 비교군 119회의 원본 예측은 실험 서버에 보존하고 전체 시도별 성적과 예측 파일 해시는 로컬 증거 묶음에 보존했다.')
table('표 A1 본문 수치의 원본 위치', ['근거','v5 디렉터리 기준 위치'],[
 ['탐지 집계와 공통 seed 비교','stopped/summary.json'],
 ['비교군 전체 학습률과 반복','runs/grid_comparison.json 및 grid_results.csv'],
 ['Laya 전체 평가 예측','runs/laya_seed20261005_predictions.npz 및 seed20261006 파일'],
 ['요청별 원본 추론 시간','speed/*_seed*.json 및 verified_summary.json'],
 ['가중치와 실행 소스 해시','speed/checkpoint_audit.json'],
 ['측정 조건과 장치 사전 상태','speed/protocol.json 및 preflight.json'],
 ['종합 보고서와 증거 묶음','STOPPED_REPORT.ko.md 및 stopped_report_and_evidence.zip'],
 ],[5.0,11.6])
note('위 경로는 실험 자료의 위치이며 공개 저장소에 업로드됐다는 뜻은 아니다. 재현용 데이터는 UQ 공식 배포에서 확보한다. 모든 속도 결과는 가중치 해시의 측정 전후 일치를 확인했다.')
p('데이터와 모델의 식별값은 다음과 같다. 전체 원본·체크포인트·소스 해시는 위 실험 자료에서 추적할 수 있다.')
for label,value in [
    ('NF-UNSW-NB15-v2 CSV SHA-256','05019e3ac8d55b3f074f9f042623bbbb1af2527b8b78558c69353ccb2ccefa3f'),
    ('Laya 기본 모델 리비전','7b928d828b7b0e022f929d9bd2e44165aa270148'),
    ('FlowTransformer 코드 커밋','52236c8c9feccdeb44acf578789db7b8b36bacea'),
    ('추론 측정 프로토콜 SHA-256',sha(ROOT/'speed/protocol.json')),
]:
    q=note(label+'\n'+value)
    for r in q.runs:r.font.name='Consolas';r.font.size=Pt(8)

doc.core_properties.title='Laya 아키텍처 기반 저지연 언어모델 위협탐지와 자동 대응 설계'
doc.core_properties.subject='공개 네트워크 흐름 실험과 저지연 의사결정 구조'
doc.core_properties.author=''
doc.core_properties.keywords='Laya, network intrusion detection, decision model, latency, correlation'
OUTPUT=HERE/'Laya_Threat_Detection_Paper.ko.docx'
doc.save(OUTPUT)
sources=['stopped/summary.json','runs/grid_comparison.json','speed/verified_summary.json','speed/protocol.json','speed/checkpoint_audit.json','train_laya.py','laya_adapter.py','common.py']
manifest=dict(created='2026-10-07',document=OUTPUT.name,sha256=sha(OUTPUT),
    source_hashes={f:sha(ROOT/f) for f in sources},references=[dict(number=i,text=t,url=u) for i,(t,u) in enumerate(REFERENCES,1)],
    claims=[
        dict(claim='Laya local warm p95 below 200ms',status='measured',source='speed/verified_summary.json'),
        dict(claim='Fast decision plus asynchronous explanation is a promising architectural direction',status='proposal and interpretation',source='design and measured inference feasibility'),
        dict(claim='Language and correlation reasoning preserved after fine-tuning',status='not established'),
        dict(claim='Laya beats selected FlowTransformer baseline',status='not supported for F1 FPR latency'),
        dict(claim='Jev or Laya is faster than a generative LLM on this detection task',status='not directly measured'),
        dict(claim='Automatic actions reduce attack spread',status='proposed not evaluated'),
    ])
(HERE/'paper_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(dict(output=str(OUTPUT),paragraphs=len(doc.paragraphs),tables=len(doc.tables),sources=len(REFERENCES),sha256=manifest['sha256']),ensure_ascii=False))
